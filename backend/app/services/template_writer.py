"""
Template Writer Service

Generates Excel (XLSX) files from mapping results using schema column definitions.
Handles parent-child variations and maintains proper column ordering.
"""
from typing import Any, Dict, List, Optional
from pathlib import Path
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# Styling constants
HEADER_FONT = Font(bold=True, size=11, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="4A5568", end_color="4A5568", fill_type="solid")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)

REQUIRED_HEADER_FILL = PatternFill(start_color="C53030", end_color="C53030", fill_type="solid")

DATA_ALIGNMENT = Alignment(horizontal="left", vertical="top", wrap_text=True)

THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

# Confidence-based cell colors
CONFIDENCE_COLORS = {
    "high": PatternFill(start_color="C6F6D5", end_color="C6F6D5", fill_type="solid"),  # Green
    "medium": PatternFill(start_color="FEFCBF", end_color="FEFCBF", fill_type="solid"),  # Yellow
    "low": PatternFill(start_color="FED7D7", end_color="FED7D7", fill_type="solid"),  # Red
}


def get_confidence_fill(confidence: float) -> Optional[PatternFill]:
    """Get cell fill color based on confidence level."""
    if confidence >= 0.8:
        return CONFIDENCE_COLORS["high"]
    elif confidence >= 0.5:
        return CONFIDENCE_COLORS["medium"]
    elif confidence > 0:
        return CONFIDENCE_COLORS["low"]
    return None


def generate_xlsx(
    mapping: Dict[str, Dict[str, Any]],
    schema: Dict[str, Any],
    output_path: str,
    include_confidence_colors: bool = True,
    include_header_row: bool = True,
) -> str:
    """
    Generate an Excel file from mapping results.

    Args:
        mapping: Mapping results {field: {value, source, confidence}}
        schema: Template schema with columns definition
        output_path: Path to save the Excel file
        include_confidence_colors: Whether to color cells by confidence
        include_header_row: Whether to include header row

    Returns:
        Path to the generated file
    """
    wb = Workbook()
    ws = wb.active
    ws.title = f"{schema.get('marketplace', 'export')}_{schema.get('category', 'data')}"

    columns = schema.get("columns", [])

    # Sort columns by col_index
    sorted_columns = sorted(columns, key=lambda c: c.get("col_index", 0))

    # Write header row
    if include_header_row:
        for col in sorted_columns:
            col_idx = col.get("col_index", 0) + 1  # Excel is 1-indexed
            cell = ws.cell(row=1, column=col_idx, value=col.get("header", col["canonical"]))

            cell.font = HEADER_FONT
            cell.alignment = HEADER_ALIGNMENT
            cell.border = THIN_BORDER

            # Different color for required fields
            if col.get("required", False):
                cell.fill = REQUIRED_HEADER_FILL
            else:
                cell.fill = HEADER_FILL

    # Determine start row for data
    data_start_row = 2 if include_header_row else 1

    # Check if schema has variations
    has_variations = schema.get("has_variations", False)

    if has_variations:
        # Write parent row and child rows
        _write_variation_rows(
            ws, mapping, sorted_columns, data_start_row, include_confidence_colors
        )
    else:
        # Write single row
        _write_single_row(
            ws, mapping, sorted_columns, data_start_row, include_confidence_colors
        )

    # Adjust column widths
    _auto_adjust_column_widths(ws, sorted_columns)

    # Save file
    output_path_obj = Path(output_path)
    output_path_obj.parent.mkdir(parents=True, exist_ok=True)

    wb.save(output_path)
    wb.close()

    return str(output_path)


def _write_single_row(
    ws,
    mapping: Dict[str, Dict[str, Any]],
    columns: List[Dict[str, Any]],
    row: int,
    include_confidence_colors: bool,
) -> None:
    """Write a single data row."""
    for col in columns:
        col_idx = col.get("col_index", 0) + 1
        canonical = col["canonical"]

        field_mapping = mapping.get(canonical, {})
        value = field_mapping.get("value", "")
        confidence = field_mapping.get("confidence", 0.0)

        cell = ws.cell(row=row, column=col_idx, value=value)
        cell.alignment = DATA_ALIGNMENT
        cell.border = THIN_BORDER

        if include_confidence_colors:
            fill = get_confidence_fill(confidence)
            if fill:
                cell.fill = fill


def _write_variation_rows(
    ws,
    mapping: Dict[str, Dict[str, Any]],
    columns: List[Dict[str, Any]],
    start_row: int,
    include_confidence_colors: bool,
) -> int:
    """
    Write parent and child variation rows.

    Returns:
        Number of rows written
    """
    # Check for variants in mapping
    variants = mapping.get("_variants", [])

    if not variants:
        # No variants - write single row
        _write_single_row(ws, mapping, columns, start_row, include_confidence_colors)
        return 1

    rows_written = 0

    # Write parent row first
    parent_mapping = {k: v for k, v in mapping.items() if not k.startswith("_")}
    _write_single_row(ws, parent_mapping, columns, start_row, include_confidence_colors)
    rows_written += 1

    # Write child rows for each variant
    for variant in variants:
        row = start_row + rows_written

        # Merge parent data with variant data
        variant_mapping = parent_mapping.copy()
        variant_mapping.update(variant)

        _write_single_row(ws, variant_mapping, columns, row, include_confidence_colors)
        rows_written += 1

    return rows_written


def _auto_adjust_column_widths(
    ws,
    columns: List[Dict[str, Any]],
    min_width: int = 10,
    max_width: int = 50,
) -> None:
    """Auto-adjust column widths based on content."""
    for col in columns:
        col_idx = col.get("col_index", 0) + 1
        col_letter = get_column_letter(col_idx)

        # Get header length
        header_len = len(col.get("header", col["canonical"]))

        # Get max content length in column
        max_len = header_len
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value:
                    cell_len = len(str(cell.value))
                    max_len = max(max_len, cell_len)

        # Set width (with padding)
        width = min(max(max_len + 2, min_width), max_width)
        ws.column_dimensions[col_letter].width = width


def generate_xlsx_batch(
    mappings: List[Dict[str, Any]],
    schema: Dict[str, Any],
    output_path: str,
) -> str:
    """
    Generate an Excel file with multiple products.

    Args:
        mappings: List of mapping results
        schema: Template schema
        output_path: Path to save the Excel file

    Returns:
        Path to the generated file
    """
    wb = Workbook()
    ws = wb.active
    ws.title = f"{schema.get('marketplace', 'export')}_{schema.get('category', 'data')}"

    columns = schema.get("columns", [])
    sorted_columns = sorted(columns, key=lambda c: c.get("col_index", 0))

    # Write header row
    for col in sorted_columns:
        col_idx = col.get("col_index", 0) + 1
        cell = ws.cell(row=1, column=col_idx, value=col.get("header", col["canonical"]))
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGNMENT
        cell.border = THIN_BORDER

        if col.get("required", False):
            cell.fill = REQUIRED_HEADER_FILL
        else:
            cell.fill = HEADER_FILL

    # Write data rows
    current_row = 2
    for mapping in mappings:
        has_variations = schema.get("has_variations", False)

        if has_variations:
            rows_written = _write_variation_rows(
                ws, mapping, sorted_columns, current_row, True
            )
            current_row += rows_written
        else:
            _write_single_row(ws, mapping, sorted_columns, current_row, True)
            current_row += 1

    # Adjust column widths
    _auto_adjust_column_widths(ws, sorted_columns)

    # Save file
    output_path_obj = Path(output_path)
    output_path_obj.parent.mkdir(parents=True, exist_ok=True)

    wb.save(output_path)
    wb.close()

    return str(output_path)


def generate_template_preview(
    mapping: Dict[str, Dict[str, Any]],
    schema: Dict[str, Any],
) -> List[Dict[str, Any]]:
    """
    Generate a preview of the template data (without creating file).

    Args:
        mapping: Mapping results
        schema: Template schema

    Returns:
        List of row dictionaries for preview
    """
    columns = schema.get("columns", [])
    sorted_columns = sorted(columns, key=lambda c: c.get("col_index", 0))

    rows = []

    # Header row
    header_row = {
        "_type": "header",
        "_columns": [
            {
                "col_index": col.get("col_index", 0),
                "header": col.get("header", col["canonical"]),
                "canonical": col["canonical"],
                "required": col.get("required", False),
            }
            for col in sorted_columns
        ],
    }
    rows.append(header_row)

    # Data row
    data_row = {
        "_type": "data",
        "_values": {},
    }

    for col in sorted_columns:
        canonical = col["canonical"]
        field_mapping = mapping.get(canonical, {})

        data_row["_values"][canonical] = {
            "col_index": col.get("col_index", 0),
            "value": field_mapping.get("value"),
            "source": field_mapping.get("source", "not_found"),
            "confidence": field_mapping.get("confidence", 0.0),
        }

    rows.append(data_row)

    return rows
