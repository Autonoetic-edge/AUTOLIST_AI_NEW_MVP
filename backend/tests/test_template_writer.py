"""
Tests for the template writer service.
"""
import os
import tempfile
import pytest
from pathlib import Path

from openpyxl import load_workbook

from app.services.template_writer import (
    generate_xlsx,
    generate_xlsx_batch,
    generate_template_preview,
    get_confidence_fill,
)
from app.services.validator_service import (
    validate_mapping,
    validate_for_export,
    ValidationResult,
)


# Sample schema for testing
SAMPLE_SCHEMA = {
    "schema_id": "test_shirt_v1",
    "marketplace": "amazon",
    "category": "shirt",
    "has_variations": False,
    "columns": [
        {"col_index": 0, "header": "SKU", "canonical": "sku", "type": "string", "required": True},
        {"col_index": 1, "header": "Title", "canonical": "title", "type": "string", "required": True},
        {"col_index": 2, "header": "Brand", "canonical": "brand", "type": "string", "required": True},
        {"col_index": 3, "header": "Price", "canonical": "price", "type": "float", "required": True},
        {"col_index": 4, "header": "Color", "canonical": "color", "type": "string", "required": False},
    ],
}


# Sample mapping for testing
SAMPLE_MAPPING = {
    "sku": {"value": "TEST-001", "source": "variant.sku", "confidence": 0.95},
    "title": {"value": "Blue Cotton T-Shirt", "source": "product.title", "confidence": 0.95},
    "brand": {"value": "TestBrand", "source": "product.vendor", "confidence": 0.90},
    "price": {"value": "29.99", "source": "variant.price", "confidence": 0.95},
    "color": {"value": "Blue", "source": "extracted.colors", "confidence": 0.70},
}


class TestValidatorService:
    """Tests for the validator service."""

    def test_validate_valid_mapping(self):
        """Test validation of a valid mapping."""
        result = validate_mapping(SAMPLE_MAPPING, SAMPLE_SCHEMA)
        assert result.valid is True
        assert len(result.errors) == 0

    def test_validate_missing_required(self):
        """Test validation with missing required field."""
        mapping = SAMPLE_MAPPING.copy()
        mapping["brand"] = {"value": None, "source": "not_found", "confidence": 0.0}

        result = validate_mapping(mapping, SAMPLE_SCHEMA)
        assert result.valid is False
        assert any(e.field == "brand" for e in result.errors)

    def test_validate_invalid_enum(self):
        """Test validation with invalid enum value."""
        schema = SAMPLE_SCHEMA.copy()
        schema["columns"] = list(SAMPLE_SCHEMA["columns"])
        schema["columns"].append({
            "col_index": 5,
            "header": "Size",
            "canonical": "size",
            "type": "string",
            "required": False,
            "enum": ["S", "M", "L", "XL"],
        })

        mapping = SAMPLE_MAPPING.copy()
        mapping["size"] = {"value": "XXL", "source": "variant", "confidence": 0.80}

        result = validate_mapping(mapping, schema)
        assert result.valid is False
        assert any("enum" in e.message.lower() or "allowed" in e.message.lower() for e in result.errors)

    def test_validate_for_export_low_confidence(self):
        """Test export validation with low confidence."""
        mapping = SAMPLE_MAPPING.copy()
        mapping["brand"] = {"value": "TestBrand", "source": "ai", "confidence": 0.40}

        result = validate_for_export(mapping, SAMPLE_SCHEMA, confidence_threshold=0.5)
        assert result.valid is False
        assert any("confidence" in e.message.lower() for e in result.errors)


class TestTemplateWriter:
    """Tests for the template writer."""

    def test_generate_xlsx_creates_file(self):
        """Test that XLSX file is created."""
        with tempfile.TemporaryDirectory() as tmpdir:
            output_path = os.path.join(tmpdir, "test_output.xlsx")

            result_path = generate_xlsx(SAMPLE_MAPPING, SAMPLE_SCHEMA, output_path)

            assert os.path.exists(result_path)
            assert result_path == output_path

    def test_generate_xlsx_has_headers(self):
        """Test that generated file has correct headers."""
        with tempfile.TemporaryDirectory() as tmpdir:
            output_path = os.path.join(tmpdir, "test_output.xlsx")
            generate_xlsx(SAMPLE_MAPPING, SAMPLE_SCHEMA, output_path)

            wb = load_workbook(output_path)
            ws = wb.active

            # Check headers match schema order
            expected_headers = ["SKU", "Title", "Brand", "Price", "Color"]
            for i, expected in enumerate(expected_headers, start=1):
                assert ws.cell(row=1, column=i).value == expected

            wb.close()

    def test_generate_xlsx_has_data(self):
        """Test that generated file has correct data."""
        with tempfile.TemporaryDirectory() as tmpdir:
            output_path = os.path.join(tmpdir, "test_output.xlsx")
            generate_xlsx(SAMPLE_MAPPING, SAMPLE_SCHEMA, output_path)

            wb = load_workbook(output_path)
            ws = wb.active

            # Check data row (row 2)
            assert ws.cell(row=2, column=1).value == "TEST-001"  # SKU
            assert ws.cell(row=2, column=2).value == "Blue Cotton T-Shirt"  # Title
            assert ws.cell(row=2, column=3).value == "TestBrand"  # Brand
            assert ws.cell(row=2, column=4).value == "29.99"  # Price
            assert ws.cell(row=2, column=5).value == "Blue"  # Color

            wb.close()

    def test_generate_xlsx_column_order(self):
        """Test that columns are in correct order by col_index."""
        schema = {
            "schema_id": "test",
            "columns": [
                {"col_index": 2, "header": "Third", "canonical": "third"},
                {"col_index": 0, "header": "First", "canonical": "first"},
                {"col_index": 1, "header": "Second", "canonical": "second"},
            ],
        }
        mapping = {
            "first": {"value": "1", "source": "test", "confidence": 0.9},
            "second": {"value": "2", "source": "test", "confidence": 0.9},
            "third": {"value": "3", "source": "test", "confidence": 0.9},
        }

        with tempfile.TemporaryDirectory() as tmpdir:
            output_path = os.path.join(tmpdir, "test_output.xlsx")
            generate_xlsx(mapping, schema, output_path)

            wb = load_workbook(output_path)
            ws = wb.active

            # Headers should be sorted by col_index
            assert ws.cell(row=1, column=1).value == "First"
            assert ws.cell(row=1, column=2).value == "Second"
            assert ws.cell(row=1, column=3).value == "Third"

            wb.close()


class TestBatchGeneration:
    """Tests for batch XLSX generation."""

    def test_generate_xlsx_batch(self):
        """Test batch generation with multiple products."""
        mappings = [
            {
                "sku": {"value": "PROD-001", "source": "variant", "confidence": 0.95},
                "title": {"value": "Product 1", "source": "product", "confidence": 0.95},
                "brand": {"value": "Brand A", "source": "product", "confidence": 0.90},
                "price": {"value": "19.99", "source": "variant", "confidence": 0.95},
                "color": {"value": "Red", "source": "extracted", "confidence": 0.70},
            },
            {
                "sku": {"value": "PROD-002", "source": "variant", "confidence": 0.95},
                "title": {"value": "Product 2", "source": "product", "confidence": 0.95},
                "brand": {"value": "Brand B", "source": "product", "confidence": 0.90},
                "price": {"value": "24.99", "source": "variant", "confidence": 0.95},
                "color": {"value": "Blue", "source": "extracted", "confidence": 0.80},
            },
        ]

        with tempfile.TemporaryDirectory() as tmpdir:
            output_path = os.path.join(tmpdir, "batch_output.xlsx")
            generate_xlsx_batch(mappings, SAMPLE_SCHEMA, output_path)

            wb = load_workbook(output_path)
            ws = wb.active

            # Should have header + 2 data rows
            assert ws.max_row == 3

            # Check first product
            assert ws.cell(row=2, column=1).value == "PROD-001"

            # Check second product
            assert ws.cell(row=3, column=1).value == "PROD-002"

            wb.close()


class TestTemplatePreview:
    """Tests for template preview generation."""

    def test_generate_preview(self):
        """Test preview generation."""
        preview = generate_template_preview(SAMPLE_MAPPING, SAMPLE_SCHEMA)

        assert len(preview) == 2

        # Check header row
        header = preview[0]
        assert header["_type"] == "header"
        assert len(header["_columns"]) == 5

        # Check data row
        data = preview[1]
        assert data["_type"] == "data"
        assert "sku" in data["_values"]
        assert data["_values"]["sku"]["value"] == "TEST-001"
        assert data["_values"]["sku"]["confidence"] == 0.95


class TestConfidenceColors:
    """Tests for confidence-based coloring."""

    def test_high_confidence_color(self):
        """Test high confidence returns green fill."""
        fill = get_confidence_fill(0.90)
        assert fill is not None
        assert fill.start_color.rgb == "00C6F6D5"  # Green

    def test_medium_confidence_color(self):
        """Test medium confidence returns yellow fill."""
        fill = get_confidence_fill(0.60)
        assert fill is not None
        assert fill.start_color.rgb == "00FEFCBF"  # Yellow

    def test_low_confidence_color(self):
        """Test low confidence returns red fill."""
        fill = get_confidence_fill(0.30)
        assert fill is not None
        assert fill.start_color.rgb == "00FED7D7"  # Red

    def test_zero_confidence_no_color(self):
        """Test zero confidence returns no fill."""
        fill = get_confidence_fill(0.0)
        assert fill is None
